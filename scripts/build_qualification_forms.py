from pathlib import Path
from itertools import count
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
import argparse
_parser=argparse.ArgumentParser()
_parser.add_argument('--output-dir', type=Path, default=Path(__file__).resolve().parents[1]/'brief')
OUT=_parser.parse_args().output_dir
OUT.mkdir(parents=True, exist_ok=True)
IDS=count(1000)

def addfield(p,tag,placeholder):
    sdt=OxmlElement('w:sdt');pr=OxmlElement('w:sdtPr')
    for key,val in [('alias',tag),('tag',tag),('id',str(next(IDS)))]:
        e=OxmlElement('w:'+key);e.set(qn('w:val'),val);pr.append(e)
    text=OxmlElement('w:text');text.set(qn('w:multiLine'),'1');pr.append(text)
    content=OxmlElement('w:sdtContent');r=OxmlElement('w:r');rp=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),'777777');rp.append(color);shd=OxmlElement('w:shd');shd.set(qn('w:fill'),'F2F2F2');rp.append(shd);r.append(rp);t=OxmlElement('w:t');t.text=placeholder;r.append(t);content.append(r);sdt.append(pr);sdt.append(content);p._p.append(sdt)

def form(lang):
    zh=lang=='zh';D=Document();sec=D.sections[0];sec.page_width=Inches(8.27);sec.page_height=Inches(11.69);sec.top_margin=sec.bottom_margin=Inches(.7);sec.left_margin=sec.right_margin=Inches(.75)
    normal=D.styles['Normal'];normal.font.name='Arial Unicode MS' if zh else 'Arial';normal.font.size=Pt(10.5);normal.font.color.rgb=RGBColor(0,0,0);normal.paragraph_format.space_after=Pt(6);normal.paragraph_format.line_spacing=1.15
    for sty,size in [('Title',23),('Subtitle',11),('Heading 1',16),('Heading 2',12),('Heading 3',11)]:
        st=D.styles[sty];st.font.name='Arial Unicode MS' if zh else 'Arial';st.font.italic=False;st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0);st.paragraph_format.space_before=Pt(10);st.paragraph_format.space_after=Pt(7)
    for st in D.styles:
        if st.type==1:
            rpr=st.element.get_or_add_rPr();rf=rpr.find(qn('w:rFonts'))
            if rf is None:rf=OxmlElement('w:rFonts');rpr.insert(0,rf)
            rf.set(qn('w:eastAsia'),'Arial Unicode MS')
            if zh:
                rf.set(qn('w:ascii'),'Arial Unicode MS');rf.set(qn('w:hAnsi'),'Arial Unicode MS')
            for borders in list(st.element.iter(qn('w:pBdr'))):borders.getparent().remove(borders)
    D.core_properties.author='XCITY';D.core_properties.title='XCITY 资格申请工作表及材料清单（可选）' if zh else 'XCITY Qualification Application Worksheet and Materials Checklist (Optional)'
    def P(c,e,style=None):return D.add_paragraph(c if zh else e,style)
    def H(c,e,level=1):return D.add_heading(c if zh else e,level)
    def F(c,e,tag,multiline=False):
        p=D.add_paragraph();p.paragraph_format.keep_with_next=multiline;p.add_run((c if zh else e)+'  ').bold=True
        addfield(p,tag,'[点击填写]' if zh else '[Click to complete]')
        if multiline:p.paragraph_format.space_after=Pt(12)
        return p
    def note(c,e):
        p=P(c,e);p.paragraph_format.space_after=Pt(8)
        for r in p.runs:r.font.size=Pt(9)
    def table(headers,rows,widths):
        tb=D.add_table(rows=1,cols=len(headers));tb.alignment=WD_TABLE_ALIGNMENT.CENTER;tb.autofit=False
        for column,width in zip(tb.columns,widths):column.width=Inches(width)
        for cell,width,header in zip(tb.rows[0].cells,widths,headers):
            cell.width=Inches(width);cell.text=header;pr=cell._tc.get_or_add_tcPr();sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'EDEDED');pr.append(sh)
            for run in cell.paragraphs[0].runs:run.bold=True
        trpr=tb.rows[0]._tr.get_or_add_trPr();rep=OxmlElement('w:tblHeader');trpr.append(rep)
        for idx,row in enumerate(rows):
            cells=tb.add_row().cells
            for j,(cell,width,value) in enumerate(zip(cells,widths,row)):
                cell.width=Inches(width);p=cell.paragraphs[0]
                if value is None:addfield(p,f'{lang}_TABLE_{next(IDS)}','[填写]' if zh else '[Complete]')
                else:p.add_run(value)
                cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=tb.rows[-1]._tr.get_or_add_trPr();nr=OxmlElement('w:cantSplit');pr.append(nr)
        for row in tb.rows:
            for cell in row.cells:
                pr=cell._tc.get_or_add_tcPr();marg=OxmlElement('w:tcMar')
                for k in ['top','left','bottom','right']:
                    el=OxmlElement('w:'+k);el.set(qn('w:w'),'90');el.set(qn('w:type'),'dxa');marg.append(el)
                pr.append(marg)
                for p in cell.paragraphs:p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(2);p.paragraph_format.line_spacing=1.05
        return tb
    def page(c,e):D.add_page_break();H(c,e)
    # Page 1
    P('XCITY 设计机构征集','XCITY Design Practice Call','Subtitle')
    P('资格申请工作表及材料清单','Application Worksheet\nand Materials Checklist','Title')
    P('本工作表仅供可选准备。官网简短注册后，我们将邮件发送专属申请链接；请在官网逐项填写资格表、提供适用证明与团队业绩、上传1–10页构想PDF，可保存并最终提交。无需上传Word，邮件不收申请附件。','This optional worksheet helps you prepare. Register on the official website; we will email a private application link. Complete the online form, evidence and references, upload a 1–10-page ideas PDF, save progress and finally submit. No Word upload is required. Do not email application attachments.')
    note('注册截止：2026年12月31日23:59 UTC。完整在线申请截止：2027年1月31日23:59 UTC。2027年2月遴选三家。','Registration deadline: 31 December 2026, 23:59 UTC. Complete online application: 31 January 2027, 23:59 UTC. Three teams will be selected in February 2027.')
    note('可使用中文或英文。灰色字段可整理草稿，最终请在专属网页填写；不适用项请说明，无需准备不适用的证书。本Word不构成正式提交。','Chinese or English is accepted. Use grey fields for preparation, then complete the private online form. Explain inapplicable items; do not obtain irrelevant certificates. This Word file is not a formal submission.')
    H('1 机构与联系人','1 Institution and contact')
    for c,e,t in [
      ('申请编号 如已收到','Application reference if received','REF'),('机构法定名称','Institution legal name','ORG'),('注册国家或地区','Country or jurisdiction of registration','COUNTRY'),('登记地址','Registered address','ADDRESS'),('申请方式 单一机构或联合体','Applicant type Single practice or consortium','TYPE'),('牵头机构 联合体填写','Lead entity Consortium only','LEAD'),('联系人姓名与职务','Contact name and position','CONTACT'),('联系邮箱','Email','EMAIL'),('电话 含国家代码 如有','Telephone with country code if available','PHONE'),('机构网站 如有','Practice website if available','WEB')]:F(c,e,t)
    note('如未收到申请编号，可留空。请使用与官网在线申请相同的机构名称和邮箱，便于核对。','Leave the reference blank if none was received. Use the same institution name and email as the online application so we can match your materials.')
    # Page 2
    page('2 注册与专业资格','2 Registration and professional qualification')
    P('提供注册证书、营业执照或官方登记查询链接，任选一种可核实依据即可。接受普通电子版或扫描件，不要求公证或纸质原件。','Provide one verifiable basis: a registration certificate, business record or official registry link. Ordinary electronic copies or scans are accepted. Notarisation and paper originals are not requested.')
    F('注册机构或登记机关','Registration authority','AUTH')
    F('注册或营业登记编号 如适用','Registration or business record number if applicable','REGNO')
    F('证明文件名或官方查询链接','Evidence filename or official registry URL','REGEVID')
    H('适用的专业资格','Applicable professional qualification',2)
    P('说明注册所在地适用的规划、建筑或景观设计执业或经营制度，以及本机构或联合体如何满足。所在地不实行机构资质分级的，可说明制度并提供等效依据，不统一要求某国资质等级。','Explain the planning, architecture or landscape licensing or business-registration arrangements that apply in your jurisdiction and how the practice or consortium meets them. Where firm-level grading does not exist, explain the system and equivalent evidence. No single country’s qualification grade is universally required.')
    F('制度与适用情况说明','Applicable arrangements and explanation','LICENCE_NOTE',True)
    F('专业资格名称 持有人及有效期 如适用','Qualification holder and validity if applicable','LICENCE')
    F('证明文件名或官方核验链接','Evidence filename or official verification URL','LICENCE_EVID')
    H('联合体成员 仅联合体填写','Consortium members Consortium applicants only',2)
    note('单一机构可填不适用。联合体请附已签署的协议或成员及职责说明，明确牵头方；成员不得再单独或加入其他团队重复申请。','Single practices may mark this section not applicable. Consortia must attach a signed agreement or signed member-and-role statement identifying the lead. A member may not also apply independently or join another team.')
    table(['成员机构' if zh else 'Member entity','国家或地区' if zh else 'Jurisdiction','职责与牵头方' if zh else 'Role and lead'],[[None,None,None] for _ in range(3)],[2.2,1.5,3.07])
    F('联合体文件名 如适用','Consortium evidence filename if applicable','CONSORTIUM')
    # Page 3
    page('3 实际参赛团队与专业分工','3 Actual team and responsibilities')
    P('请列出实际参加本项目的主创与主要成员，以及各自的职责。可按实际团队增行。','Identify the principal designer and the key people who will actually work on this project, with their responsibilities. Add rows as needed for the actual team.')
    table(['姓名' if zh else 'Name','机构与职务' if zh else 'Entity and role','专业或适用资格' if zh else 'Discipline or qualification','本项目职责' if zh else 'Project responsibility'],[[None,None,None,None] for _ in range(5)],[1.15,1.8,1.72,2.1])
    H('团队能力与协作安排','Team coverage and collaboration',2)
    P('说明团队如何覆盖总体规划、城市设计、景观与生态，以及水资源、能源、交通和实施等专业协作。专业人员尚未确定的，请说明安排，不要填写未确认的参与者。','Explain coverage of masterplanning, urban design, landscape and ecology, and specialist collaboration for water, energy, transport and delivery. If a specialist is not yet confirmed, describe the arrangement rather than naming an unconfirmed participant.')
    F('总体规划与城市设计','Masterplanning and urban design','TEAM_URBAN',True)
    F('景观与生态','Landscape and ecology','TEAM_LANDSCAPE',True)
    F('水资源 能源 交通与实施协作','Water energy transport and delivery collaboration','TEAM_SPECIALISTS',True)
    F('主创姓名与联系邮箱','Principal designer name and email','PRINCIPAL')
    note('申请阶段不要求事先具备阿根廷当地合作方。按任务书，选定团队的本地协作由 XCITY 安排。','An Argentine partner is not required at application. As stated in the brief, XCITY arranges local collaboration for the selected team.')
    # Page 4
    page('4 相关项目业绩','4 Relevant project experience')
    P('选择能说明本项目相关能力的代表业绩。写明机构及主创实际承担的职责，不以作品集图片代替职责说明。可按需复制本页。','Select references that demonstrate capability relevant to this project. State the practice’s and principal designer’s actual contributions; portfolio images do not replace role descriptions. Copy this page as needed.')
    for n in [1,2]:
      H(f'项目记录 {n}',f'Project record {n}',2)
      for c,e,t in [
        ('项目名称与地点','Project name and location','NAME'),('规模与单位','Scale with units','SCALE'),('年份与状态 如建成 在建或未实施','Year and status Built ongoing or unbuilt','STATUS'),('申请机构及主创的实际职责','Actual role of applicant and principal designer','ROLE'),('与本项目的相关性','Relevance to this competition','RELEVANCE'),('作品集文件名或项目链接','Portfolio filename or project URL','LINK')]:F(c,e,f'PROJECT_{n}_{t}')
    note('涉及联合设计时请注明合作方与分工。引用的作品和图片应有权用于本次申请；不要求提供客户保密资料。','For joint work, identify collaborators and responsibilities. Ensure references and images may be used for this application; client-confidential material is not requested.')
    # Page 5
    page('5 材料清单','5 Materials checklist')
    P('可在右栏整理拟上传文件名或链接；不适用项说明原因。请在官网专属页逐项填写及上传，无需合并成一份大型PDF，也无需上传本Word。','Use the right column to prepare filenames or links. Explain conditional items that do not apply. Complete and upload each item on your private online page; no combined PDF or upload of this Word worksheet is required.')
    rows=[
      ('在线资格申请表','Online qualification application form','在专属网页填写 无需上传Word','Complete on your private page No Word upload'),
      ('机构注册依据','Legal registration evidence','证书 营业记录或官方链接 任选一项','Certificate business record or official link Choose one'),
      ('适用专业资格','Applicable professional evidence','资格文件 或当地制度与等效依据说明','Evidence or explanation of local arrangements and equivalent basis'),
      ('实际团队与分工','Actual team and roles','在线填写 可用本表第3部分准备','Complete online Section 3 is a preparation aid'),
      ('相关项目业绩','Relevant project references','在线填写 作品集文件或项目链接','Complete online Portfolio file or project links'),
      ('初步构想','Initial ideas','单独1–10页PDF 在线上传','Separate 1–10-page PDF Upload online'),
      ('联合体协议或成员职责说明','Consortium agreement or member-and-role statement','仅联合体 需签署','Consortia only Signed'),
      ('真实性及利益关系声明','Accuracy and conflicts declaration','在线确认 本表第6部分仅供准备','Confirm online Section 6 is preparation only')]
    table(['材料' if zh else 'Material','可接受形式' if zh else 'Accepted form','文件名或链接' if zh else 'Filename or link'],[[(c if zh else e),(c2 if zh else e2),None] for c,e,c2,e2 in rows],[1.7,2.65,2.42])
    H('初步构想应说明什么','What the initial ideas should explain',2)
    P('以1–10页PDF说明场地判断、城市与五大系统的初步方向、首期启动逻辑、代表性空间想法及工作方法。草图、示意图或剖面即可；此阶段不要求完整总图或效果图。','Use a 1–10-page PDF to explain the site reading, initial city and five-system direction, phase-one logic, a representative spatial idea and your approach. Sketches, diagrams or sections are sufficient; a full masterplan or rendering is not required at this stage.')
    note('证明支持普通PDF、PNG或JPG；每个文件最多10 MB，全部上传合计最多25 MB。接受的证明及项目链接应可由授权审核人员访问；访问限制请在在线说明中注明。邮件仅发送专属链接，不接收申请附件。','Evidence uploads: PDF, PNG or JPG; each file up to 10 MB, total uploads up to 25 MB. Accepted evidence/project links must be accessible to authorised reviewers; explain access restrictions online. Email delivers the private link, not application attachments.')
    # Page 6
    page('6 声明与在线提交','6 Declaration and online submission')
    P('请由有权代表申请机构或联合体的人员阅读。可先在本工作表整理姓名、日期与声明，正式确认请在官网专属申请页完成。','An authorised representative should read these declarations. Prepare names, dates and statements here if useful; make the formal confirmation on the private online application page.')
    declarations=[
      ('所填信息及所附材料真实、可核实，项目署名与实际职责已如实说明。','The information and evidence are accurate and verifiable. Project credits and actual responsibilities are stated truthfully.'),
      ('列出的参赛团队及协作安排已如实说明，未将未确认人员列作已确认成员。','The stated team and collaboration arrangements are accurate. Unconfirmed participants are not presented as confirmed.'),
      ('本机构仅参加一个申请团队；联合体成员未同时单独申请或加入其他团队。','The practice participates in only one applicant team. Consortium members do not also apply independently or join another team.'),
      ('参赛人员不属于主办方工作人员、评审团成员及其直系亲属；已在下方披露其他相关关系及利益，无则写无。','Participants are not organiser staff, jury members or their immediate families. Other relevant relationships or interests are disclosed below. Write none if none.'),
      ('我同意 XCITY 私下保存并由授权人员评审本次申请材料。此确认不构成另行授予作品使用或转让权利。','I agree to private storage and review of these materials by authorised XCITY personnel. This confirmation does not grant additional rights to use or transfer the work.')]
    for i,(c,e) in enumerate(declarations,1):
      p=D.add_paragraph();p.add_run(('确认 ' if zh else 'Confirm ')+str(i)+'  ').bold=True;addfield(p,f'CONFIRM_{i}','[是／否及说明]' if zh else '[Yes or no with explanation]');D.add_paragraph(c if zh else e)
    F('利益关系披露或其他说明','Conflicts or other explanation','CONFLICTS',True)
    F('授权代表姓名与职务','Authorised representative name and position','SIGNATORY')
    F('代表的机构及确认日期','Represented entity and date of confirmation','SIGN_DATE')
    H('如何完成在线申请','How to complete the online application',2)
    P('打开主办方邮件中的专属申请链接，在官网逐项填写并提供材料。可保存后继续；核对完整性后最终提交。专属链接请妥善保存，不要公开。遇到链接或上传问题，请联系 competition@xcity.ai；请勿发送申请附件。','Open the private link emailed by the organiser. Complete each item and provide the materials on the official website; save and return as needed, then make the final submission. Keep the link private. For link or upload assistance contact competition@xcity.ai; do not email application attachments.')
    note('完整在线申请截止：2027年1月31日23:59 UTC。仅注册或保存草稿不代表完成第一阶段申请；最终提交后进入审核，但不代表资格通过或入围。','Final online submission is due by 31 January 2027, 23:59 UTC. Registration or a saved draft alone does not complete stage one. Final submission enters review; it does not mean eligibility approval or shortlisting.')
    # Running identifier and pages only
    header=sec.header.paragraphs[0];header.text='XCITY  |  '+('设计机构征集' if zh else 'DESIGN PRACTICE CALL');header.style='Caption'
    for run in header.runs:run.font.color.rgb=RGBColor(0,0,0);run.font.size=Pt(8)
    foot=sec.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    foot.add_run(('2026年10月26日  |  第 ' if zh else '26 October 2026  |  Page '))
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');foot._p.append(field)
    for run in foot.runs:run.font.size=Pt(8)
    path=OUT/('qualification-application.zh.docx' if zh else 'qualification-application.en.docx');D.save(path);print(path)
for lang in ['zh','en']:form(lang)
